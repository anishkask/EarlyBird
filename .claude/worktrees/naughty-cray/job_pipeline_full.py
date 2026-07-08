#!/usr/bin/env python3
"""
================================================================
  Job Search Pipeline
  Strategy: Find postings before 100 people apply.
================================================================
See full docstring in script for setup and usage.
Run: python job_pipeline_full.py
     python job_pipeline_full.py --no-email
     python job_pipeline_full.py --hours 12
"""

import os
import sys
import re
import time
import random
import json
import csv
import argparse
import warnings
from datetime import datetime, timezone, timedelta
from pathlib import Path
import requests
from bs4 import BeautifulSoup
import pandas as pd
from dotenv import load_dotenv
warnings.filterwarnings("ignore")

# Load environment variables from .env file
load_dotenv()

YOUR_NAME      = os.getenv("YOUR_NAME", "")
YOUR_EMAIL     = os.getenv("YOUR_EMAIL", "")
YOUR_SCHOOL    = os.getenv("YOUR_SCHOOL", "")
YOUR_LINKEDIN  = os.getenv("YOUR_LINKEDIN", "")
YOUR_GITHUB    = os.getenv("YOUR_GITHUB", "")
YOUR_PORTFOLIO = os.getenv("YOUR_PORTFOLIO", "")
MY_BACKGROUND  = os.getenv("MY_BACKGROUND", "")
EMAIL_DELAY_MIN = int(os.getenv("EMAIL_DELAY_MIN", "120"))
EMAIL_DELAY_MAX = int(os.getenv("EMAIL_DELAY_MAX", "300"))

# === Dynamic ATS Discovery ===
# Bootstrap set of known companies with confirmed ATS — used as seed for discovery
# and fallback when the dynamic watchlist cannot be built.

FALLBACK_COMPANIES = [
    {"name": "Airbnb",        "website": "https://www.airbnb.com",        "ats_type": "greenhouse", "slug": "airbnb"},
    {"name": "Brex",          "website": "https://www.brex.com",          "ats_type": "greenhouse", "slug": "brex"},
    {"name": "Coinbase",      "website": "https://www.coinbase.com",      "ats_type": "lever",      "slug": "coinbase"},
    {"name": "Notion",        "website": "https://www.notion.so",         "ats_type": "greenhouse", "slug": "notion"},
    {"name": "Vercel",        "website": "https://vercel.com",            "ats_type": "lever",      "slug": "vercel"},
    {"name": "Linear",        "website": "https://linear.app",            "ats_type": "lever",      "slug": "linear"},
    {"name": "Retool",        "website": "https://retool.com",            "ats_type": "lever",      "slug": "retool"},
    {"name": "Glydways",      "website": "https://www.glydways.com",      "ats_type": "greenhouse", "slug": "glydways"},
    {"name": "Monarch Money", "website": "https://www.monarchmoney.com",  "ats_type": "greenhouse", "slug": "monarchmoney"},
    {"name": "Eulerity",      "website": "https://eulerity.com",          "ats_type": "greenhouse", "slug": "eulerity"},
]

WATCHLIST_CACHE_FILE = "watchlist_cache.json"
WATCHLIST_CACHE_MAX_DAYS = 7
CO_OUTREACH_CACHE = Path("data") / "cold_outreach_cache.json"


def _ats_slugs_from_name(name: str) -> list:
    """Generate ATS slug candidates from a company name."""
    base = re.sub(r"[^a-z0-9\s-]", "", name.lower()).strip()
    hyphenated = re.sub(r"\s+", "-", base)
    no_space   = re.sub(r"\s+", "",  base)
    first_word = base.split()[0] if base.split() else base
    return list(dict.fromkeys([hyphenated, no_space, first_word]))  # unique, ordered


def _check_greenhouse(slug: str) -> bool:
    try:
        r = requests.get(
            f"https://boards-api.greenhouse.io/v1/boards/{slug}/jobs",
            headers=HEADERS, timeout=3,
        )
        return r.status_code == 200 and bool(r.json().get("jobs"))
    except Exception:
        return False


def _check_lever(slug: str) -> bool:
    try:
        r = requests.get(
            f"https://api.lever.co/v0/postings/{slug}?mode=json",
            headers=HEADERS, timeout=3,
        )
        return r.status_code == 200 and isinstance(r.json(), list) and bool(r.json())
    except Exception:
        return False


def build_dynamic_watchlist() -> list:
    """
    Build a dynamic list of companies with active Greenhouse or Lever boards.
    Uses a 7-day cache (watchlist_cache.json) to avoid re-checking every run.
    Seeds candidates from cold_outreach_cache.json (VC portfolio discoveries)
    plus the bootstrap FALLBACK_COMPANIES list.
    Also polls YC Work at a Startup and Wellfound for additional companies.
    Returns list of dicts: {name, website, ats_type, slug}
    """
    # 1. Return cache if fresh
    if Path(WATCHLIST_CACHE_FILE).exists():
        age_days = (time.time() - Path(WATCHLIST_CACHE_FILE).stat().st_mtime) / 86400
        if age_days < WATCHLIST_CACHE_MAX_DAYS:
            try:
                with open(WATCHLIST_CACHE_FILE) as f:
                    cached = json.load(f)
                companies = cached.get("companies", [])
                if companies:
                    n_gh = sum(1 for c in companies if c.get("ats_type") == "greenhouse")
                    n_lv = sum(1 for c in companies if c.get("ats_type") == "lever")
                    print(f"Dynamic watchlist built: {n_gh} Greenhouse companies, {n_lv} Lever companies, 0 from real-time sources (cached)")
                    return companies
            except Exception as e:
                print(f"  [WARNING] Could not read watchlist cache: {e}")

    print("  Building dynamic watchlist (checking ATS endpoints)...")

    # 2. Collect candidates from cold_outreach_cache.json + fallback list
    candidates: list[dict] = []
    if CO_OUTREACH_CACHE.exists():
        try:
            with open(CO_OUTREACH_CACHE) as f:
                data = json.load(f)
            for c in data.get("companies", []):
                if c.get("name"):
                    candidates.append({"name": c["name"], "website": c.get("website", "")})
        except Exception as e:
            print(f"  [WARNING] Could not read cold_outreach_cache.json: {e}")

    for c in FALLBACK_COMPANIES:
        candidates.append({"name": c["name"], "website": c["website"]})

    # 3. Check Greenhouse and Lever for each candidate
    verified: list[dict] = []
    seen_slugs: set = set()

    for cand in candidates:
        name = cand.get("name", "")
        if not name:
            continue
        slugs = _ats_slugs_from_name(name)
        found = False
        for slug in slugs:
            if slug in seen_slugs or not slug:
                continue
            if _check_greenhouse(slug):
                seen_slugs.add(slug)
                verified.append({"name": name, "website": cand.get("website", ""), "ats_type": "greenhouse", "slug": slug})
                found = True
                break
            if _check_lever(slug):
                seen_slugs.add(slug)
                verified.append({"name": name, "website": cand.get("website", ""), "ats_type": "lever", "slug": slug})
                found = True
                break
        _ = found  # not used further

    # 4. Real-time sources
    rt_count = 0

    # YC Work at a Startup
    try:
        r = requests.get(
            "https://www.workatastartup.com/jobs?role=eng&usa_only=true",
            headers=HEADERS, timeout=10,
        )
        soup = BeautifulSoup(r.text, "html.parser")
        for link in soup.find_all("a", href=True)[:120]:
            href = str(link.get("href", ""))
            text = link.get_text(strip=True)
            if "workatastartup.com/companies/" in href and text and 2 <= len(text) <= 60:
                for slug in _ats_slugs_from_name(text)[:2]:
                    if slug in seen_slugs:
                        continue
                    if _check_greenhouse(slug):
                        seen_slugs.add(slug)
                        verified.append({"name": text, "website": href, "ats_type": "greenhouse", "slug": slug})
                        rt_count += 1
                        break
                    if _check_lever(slug):
                        seen_slugs.add(slug)
                        verified.append({"name": text, "website": href, "ats_type": "lever", "slug": slug})
                        rt_count += 1
                        break
    except Exception as e:
        print(f"  [WARNING] YC Work at a Startup scrape failed: {e}")

    # Wellfound
    try:
        r = requests.get(
            "https://wellfound.com/jobs?role=software-engineer&remote=true",
            headers=HEADERS, timeout=10,
        )
        soup = BeautifulSoup(r.text, "html.parser")
        for tag in soup.find_all(class_=re.compile("company", re.I))[:60]:
            text = tag.get_text(strip=True)
            if not text or len(text) < 2 or len(text) > 60:
                continue
            for slug in _ats_slugs_from_name(text)[:1]:
                if slug in seen_slugs:
                    continue
                if _check_greenhouse(slug):
                    seen_slugs.add(slug)
                    verified.append({"name": text, "website": "", "ats_type": "greenhouse", "slug": slug})
                    rt_count += 1
                    break
    except Exception as e:
        print(f"  [WARNING] Wellfound scrape failed: {e}")

    # 5. Fall back to bootstrap list if discovery returned nothing
    if not verified:
        print(f"  [FALLBACK] Dynamic discovery found no companies — using fallback list")
        return FALLBACK_COMPANIES

    # 6. Cache and report
    n_gh = sum(1 for c in verified if c.get("ats_type") == "greenhouse")
    n_lv = sum(1 for c in verified if c.get("ats_type") == "lever")
    try:
        with open(WATCHLIST_CACHE_FILE, "w") as f:
            json.dump({"companies": verified, "built_at": datetime.now().isoformat()}, f, indent=2)
    except Exception as e:
        print(f"  [WARNING] Could not save watchlist_cache.json: {e}")

    print(f"Dynamic watchlist built: {n_gh} Greenhouse companies, {n_lv} Lever companies, {rt_count} from real-time sources")
    return verified


# Build dynamic watchlist on startup (uses cache when available — fast path)
COMPANIES = build_dynamic_watchlist()

# Extract ATS-specific lists for backward compatibility with scrape functions
GREENHOUSE_SLUGS = [c["slug"] for c in COMPANIES if c.get("ats_type") == "greenhouse"]
LEVER_SLUGS = [c["slug"] for c in COMPANIES if c.get("ats_type") == "lever"]
ASHBY_SLUGS = []  # No Ashby companies in this setup yet

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
}

def hours_ago(dt):
    """Calculate hours elapsed since datetime dt. Returns 999 if dt is None."""
    if dt is None:
        return 999
    now = datetime.now(timezone.utc)
    if hasattr(dt, 'astimezone'):
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
    else:
        return 999
    delta = now - dt
    return int(delta.total_seconds() / 3600)

def is_intern(title):
    """Check if job title is for an internship."""
    if not title:
        return False
    t = title.lower()
    return "intern" in t or "co-op" in t or "graduate" in t


def resume_hint(title, desc=""):
    t = (title + " " + desc).lower()
    if any(k in t for k in ["rag", "llm", "ai ", "ml ", "machine learning", "genai"]):
        return "AI/ML Resume"
    if any(k in t for k in ["data pipeline", "analytics", "etl", "sql", "data engineer"]):
        return "Data Resume"
    if any(k in t for k in ["frontend", "front-end", "react", "css", "html"]):
        return "Frontend Resume"
    return "SWE Base Resume"

def score(title, desc=""):
    """Calculate match score (0-100) for a job based on title and description."""
    if not title:
        return 0
    text = (title + " " + desc).lower()
    score_val = 0
    
    # High-value keywords (25 points each)
    high_value = ["rag", "llm", "ai ", "ml ", "machine learning", "genai", "data engineer", "analytics"]
    score_val += sum(25 for keyword in high_value if keyword in text)
    
    # Medium-value keywords (15 points each)
    medium_value = ["backend", "fullstack", "frontend", "react", "python", "java", "c++", "go"]
    score_val += sum(15 for keyword in medium_value if keyword in text)
    
    # Low-value keywords (5 points each)
    low_value = ["remote", "flexible", "hybrid", "relocation", "visa"]
    score_val += sum(5 for keyword in low_value if keyword in text)
    
    return min(score_val, 100)  # Cap at 100

def job(title,company,location,url,dt,source,desc=""):
    ha = hours_ago(dt)
    return {"title":title,"company":company,"location":location,"job_url":url,
            "hours_ago":ha,"source":source,"description":desc[:400]}

def scrape_greenhouse(max_h):
    """Scrape Greenhouse ATS boards using public API."""
    out = []
    for token in GREENHOUSE_SLUGS:
        try:
            url = f"https://boards-api.greenhouse.io/v1/boards/{token}/jobs?content=true"
            r = requests.get(url, headers=HEADERS, timeout=8)
            if r.status_code == 404:
                # Token doesn't exist, skip silently
                continue
            if r.status_code != 200:
                continue
            
            for j in r.json().get("jobs", []):
                title = j.get("title", "")
                if not is_intern(title):
                    continue
                
                # Parse timestamp from updated_at or created_at
                dt = None
                for ts_field in ["updated_at", "created_at"]:
                    ts_str = j.get(ts_field, "")
                    if ts_str:
                        try:
                            dt = datetime.fromisoformat(ts_str.replace("Z", "+00:00"))
                            break
                        except Exception as e:
                            print(f"WARNING: [Greenhouse:{token}] timestamp parse error: {e}")
                            continue
                
                # Default to 72h ago if no timestamp found
                if dt is None:
                    dt = datetime.now(timezone.utc) - timedelta(hours=72)
                
                if hours_ago(dt) > max_h:
                    continue
                
                location = j.get("location", {})
                if isinstance(location, dict):
                    loc_name = location.get("name", "")
                else:
                    loc_name = str(location)
                
                out.append(job(
                    title,
                    token.replace("-", " ").title(),
                    loc_name,
                    j.get("absolute_url", ""),
                    dt,
                    "Greenhouse",
                    BeautifulSoup(j.get("content", ""), "html.parser").get_text()[:400]
                ))
        except Exception as e:
            print(f"WARNING: [Greenhouse:{token}] {e}")
    return out

def scrape_lever(max_h):
    """Scrape Lever ATS boards using public API."""
    out = []
    for slug in LEVER_SLUGS:
        try:
            url = f"https://api.lever.co/v0/postings/{slug}?mode=json"
            r = requests.get(url, headers=HEADERS, timeout=8)
            if r.status_code == 404:
                # Company doesn't have Lever board, skip silently
                continue
            if r.status_code != 200:
                continue
            
            for j in r.json():
                title = j.get("text", "")
                if not is_intern(title):
                    continue
                
                # Parse createdAt Unix timestamp (in milliseconds)
                dt = None
                created_ms = j.get("createdAt", 0)
                if created_ms and created_ms > 0:
                    try:
                        dt = datetime.fromtimestamp(created_ms / 1000.0, tz=timezone.utc)
                    except Exception as e:
                        print(f"WARNING: [Lever:{slug}] timestamp parse error: {e}")
                
                # Default to 72h ago if no timestamp found
                if dt is None:
                    dt = datetime.now(timezone.utc) - timedelta(hours=72)
                
                if hours_ago(dt) > max_h:
                    continue
                
                location = j.get("categories", {}).get("location", "")
                out.append(job(
                    title,
                    slug.replace("-", " ").title(),
                    location,
                    j.get("hostedUrl", ""),
                    dt,
                    "Lever",
                    j.get("descriptionPlain", "")[:400]
                ))
        except Exception as e:
            print(f"WARNING: [Lever:{slug}] {e}")
    return out

def scrape_ashby(max_h):
    out=[]
    for slug in ASHBY_SLUGS:
        try:
            payload={"operationName":"ApiJobBoardWithTeams",
                     "variables":{"organizationHostedJobsPageName":slug},
                     "query":"query ApiJobBoardWithTeams($organizationHostedJobsPageName:String!){jobBoard:publishedJobBoard(organizationHostedJobsPageName:$organizationHostedJobsPageName){jobPostings{id title locationName publishedDate externalLink}}}"}
            r=requests.post("https://jobs.ashbyhq.com/api/non-user-graphql?op=ApiJobBoardWithTeams",
                            json=payload,headers=HEADERS,timeout=6)
            if r.status_code != 200:
                continue
            job_postings = (r.json().get("data", {}).get("jobBoard", {}) or {}).get("jobPostings", []) or []
            for j in job_postings:
                t = j.get("title", "")
                if not is_intern(t):
                    continue
                dt = None
                try:
                    dt = datetime.fromisoformat(str(j.get("publishedDate", "")).replace("Z", "+00:00"))
                except Exception as e:
                    print(f"WARNING: [Ashby:{slug}] timestamp parse error: {e}")
                if hours_ago(dt) > max_h:
                    continue
                url = j.get("externalLink") or f"https://jobs.ashbyhq.com/{slug}/{j.get('id','')}"
                out.append(job(
                    t,
                    slug.replace("-", " ").title(),
                    j.get("locationName", ""),
                    url,
                    dt,
                    "Ashby/YC"
                ))
        except Exception as e:
            print(f"WARNING: [Ashby:{slug}] {e}")
    return out

def scrape_wellfound():
    out=[]
    try:
        r = requests.get("https://wellfound.com/jobs?role=engineer&job_type=internship",
                        headers=HEADERS, timeout=10)
        soup = BeautifulSoup(r.text, "html.parser")
        for card in soup.select("div[class*='JobListing'], div[class*='job-listing']")[:30]:
            a = card.find("a", href=True)
            co = card.find(class_=re.compile("company", re.I))
            if a and is_intern(a.text):
                dt = datetime.now(timezone.utc) - timedelta(hours=72)
                out.append(job(
                    a.text.strip(),
                    co.text.strip() if co else "",
                    "Remote",
                    "https://wellfound.com" + a["href"],
                    dt,
                    "Wellfound"
                ))
    except Exception as e:
        print(f"WARNING: [Wellfound] {e}")
    return out

def scrape_jobspy(max_h):
    """Scrape LinkedIn and Indeed using JobSpy (ZipRecruiter removed)."""
    out = []
    try:
        from jobspy import scrape_jobs
        for term in [
            "software engineering intern 2026",
            "AI engineering intern remote",
            "backend engineer intern",
            "fullstack intern"
        ]:
            try:
                df = scrape_jobs(
                    site_name=["linkedin", "indeed"],
                    search_term=term,
                    location="United States",
                    results_wanted=15,
                    hours_old=max_h,
                    country_indeed="USA",
                    is_remote=True
                )
                if df is not None and not df.empty:
                    for _, row in df.iterrows():
                        t = str(row.get("title", ""))
                        if not is_intern(t):
                            continue
                        
                        # Convert date_posted to UTC datetime
                        dt = row.get("date_posted")
                        if hasattr(dt, "to_pydatetime"):
                            dt = dt.to_pydatetime()
                        
                        # Ensure UTC timezone
                        if dt and isinstance(dt, datetime):
                            if dt.tzinfo is None:
                                # Assume UTC if no timezone info
                                dt = dt.replace(tzinfo=timezone.utc)
                        
                        # Default to 72h ago if no timestamp
                        if dt is None:
                            dt = datetime.now(timezone.utc) - timedelta(hours=72)
                        
                        if hours_ago(dt) > max_h:
                            continue
                        
                        out.append(job(
                            t,
                            str(row.get("company", "")),
                            str(row.get("location", "")),
                            str(row.get("job_url", "")),
                            dt,
                            str(row.get("site", "")).title(),
                            str(row.get("description", ""))[:400]
                        ))
            except Exception as e:
                print(f"WARNING: [JobSpy:{term!r}] {e}")
    except ImportError:
        print("  [jobspy] skipped — run: pip install python-jobspy")
    return out

def find_contacts(company,role,client):
    try:
        r=client.messages.create(model="claude-sonnet-4-20250514",max_tokens=700,
            messages=[{"role":"user","content":f"""Find 3 best LinkedIn contacts at "{company}" for "{role}" internship.
Priority: {os.getenv("YOUR_SCHOOL", "my university")} alumni > recruiters/TA > junior/mid engineers.
Return JSON array only:
[{{"name":null,"title":"","reason":"","linkedin_search_url":"https://linkedin.com/search/results/people/?keywords=...","guessed_email":null}}]"""}])
        return json.loads(re.sub(r'```json|```','',r.content[0].text).strip())
    except Exception as e:
        print(f"WARNING: [find_contacts:{company}] {e}")
        return []

def draft_messages(company,role,url,contact,client):
    name=contact.get("name") or ""
    reason=contact.get("reason","")
    school = os.getenv("YOUR_SCHOOL", "my university")
    angle=(f"I noticed we both went to {school}!" if "alumni" in reason.lower() or school.lower().split()[0] in reason.lower()
           else f"I came across your profile while researching {company}.")
    try:
        r=client.messages.create(model="claude-sonnet-4-20250514",max_tokens=700,
            messages=[{"role":"user","content":f"""Draft outreach from {YOUR_NAME} to {name or 'this person'} at {company} re: {role}.
Connection: {angle}
Background: {MY_BACKGROUND}
Job: {url}
No em-dashes. Warm, specific, professional.
Return JSON only: {{"email_subject":"","email_body":"","linkedin_message":"(under 280 chars)"}}"""}])
        return json.loads(re.sub(r'```json|```','',r.content[0].text).strip())
    except Exception as e:
        print(f"WARNING: [draft_messages:{company}] {e}")
        return {"email_subject":"","email_body":"","linkedin_message":""}

def get_gmail():
    try:
        from google.oauth2.credentials import Credentials
        from google_auth_oauthlib.flow import InstalledAppFlow
        from google.auth.transport.requests import Request
        from googleapiclient.discovery import build
        # gmail.compose covers both send and draft creation (superset of gmail.send).
        # If your existing token.json only has gmail.send, delete it and re-run once
        # to re-authorize with the broader scope.
        SCOPES=["https://www.googleapis.com/auth/gmail.compose"]
        creds=None
        if Path("token.json").exists():
            creds=Credentials.from_authorized_user_file("token.json",SCOPES)
        if not creds or not creds.valid:
            if creds and creds.expired and creds.refresh_token:
                creds.refresh(Request())
            else:
                if not Path("credentials.json").exists():
                    return None
                creds=InstalledAppFlow.from_client_secrets_file("credentials.json",SCOPES).run_local_server(port=0)
            Path("token.json").write_text(creds.to_json())
        return build("gmail","v1",credentials=creds)
    except Exception as e:
        print(f"WARNING: [Gmail setup] {e}")
        return None

def send_email(gmail,to,subject,body,dry_run=False):
    import base64
    from email.mime.text import MIMEText
    if dry_run or not gmail:
        print(f"    [draft] {to} | {subject[:45]}")
        return False
    try:
        msg = MIMEText(body)
        msg["to"] = to
        msg["from"] = YOUR_EMAIL
        msg["subject"] = subject
        gmail.users().messages().send(userId="me",
            body={"raw":base64.urlsafe_b64encode(msg.as_bytes()).decode()}).execute()
        return True
    except Exception as e:
        print(f"    Email failed: {e}")
        return False

def write_excel(jobs,outreach):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter

    ts=datetime.now().strftime("%Y-%m-%d_%H-%M")
    fname=f"job_leads_{ts}.xlsx"
    wb=Workbook()
    hfill=PatternFill("solid",fgColor="0F2940")

    # Jobs sheet
    ws = wb.active
    ws.title = "Jobs — Apply Here"
    hdrs = [
        "#", "Role", "Company", "Location", "Source", "Posted", "Apply Link",
        "Match Score", "Applied?", "Date Applied", "Status", "Notes"
    ]
    ws.append(hdrs)
    for c in ws[1]:
        c.font=Font(name="Arial",bold=True,color="FFFFFF",size=9)
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")

    for i, j in enumerate(jobs, 1):
        ha = j.get("hours_ago", 999)
        posted = f"{ha:.0f}h ago" if ha < 999 else "Unknown"
        ws.append([
            i,
            j["title"],  # Role column
            j["company"],
            j["location"],
            j["source"],
            posted,
            j["job_url"],
            score(j["title"], j.get("description", "")),
            "",  # Applied?
            "",  # Date Applied
            "Not Applied",  # Status
            ""   # Notes
        ])
        color = (
            "D5F5E3" if ha <= 6 else
            "D6EAF8" if ha <= 24 else
            "FEF9E7" if ha <= 48 else "FDFEFE"
        )
        for c in ws[i + 1]:
            c.fill = PatternFill("solid", fgColor=color)

    for i, w in enumerate([4, 35, 22, 18, 12, 14, 50, 12, 18, 14, 14, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes="A2"

    # Outreach sheet
    ws2=wb.create_sheet("Outreach — Do This")
    hdrs2=["Company","Role","Contact Name","Contact Title","Why Reach Out",
           "LinkedIn Search URL","LinkedIn Message (copy+send manually)",
           "Email Address","Email Subject","Email Body Preview","Email Sent?","LinkedIn Sent?"]
    ws2.append(hdrs2)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")
    for row in outreach:
        ws2.append([row.get("company",""),row.get("role",""),
                    row.get("name",""),row.get("title",""),row.get("reason",""),
                    row.get("linkedin_url",""),row.get("linkedin_msg",""),
                    row.get("email",""),row.get("email_subj",""),
                    row.get("email_body","")[:250]+"...",
                    row.get("email_sent",""),
                    "SEND MANUALLY via LinkedIn"])
    for i,w in enumerate([22,30,20,22,28,50,55,28,35,45,14,25],1):
        ws2.column_dimensions[get_column_letter(i)].width=w
    ws2.freeze_panes="A2"

    # Legend
    ws3=wb.create_sheet("Legend & Setup")
    for row in [
        ["COLOR KEY",""],
        ["Green","Posted < 6 hours — apply IMMEDIATELY, you may be in first 10"],
        ["Blue","Posted < 24 hours — apply today"],
        ["Yellow","Posted < 48 hours — apply this week"],
        ["White","Posted > 48 hours — still worth applying"],
        [""],
        ["OUTREACH INSTRUCTIONS",""],
        ["LinkedIn messages","Copy from Outreach tab → open LinkedIn search URL → send. Takes 30 sec each."],
        ["Emails","Sent automatically to guessed addresses. Check your Sent folder."],
        [""],
        ["SETUP",""],
        ["1. Install","pip install python-jobspy pandas openpyxl anthropic requests beautifulsoup4"],
        ["   Gmail","pip install google-auth google-auth-oauthlib google-auth-httplib2 google-api-python-client"],
        ["2. API key","export ANTHROPIC_API_KEY='sk-...'"],
        ["3. Gmail","Get credentials.json from Google Cloud Console (enable Gmail API, OAuth Desktop)"],
        ["4. Run","python job_pipeline_full.py"],
        ["   Dry run","python job_pipeline_full.py --no-email"],
        ["   Fresh only","python job_pipeline_full.py --hours 12"],
    ]:
        ws3.append(row)
    ws3.column_dimensions["A"].width=20
    ws3.column_dimensions["B"].width=80

    # NOTE: caller is responsible for wb.save(fname) so that the Cold Outreach
    # module can append its sheet before the file is written to disk.
    return fname, wb

def run_pipeline_api(
    api_key: str,
    hours: int = 72,
    cold_outreach_enabled: bool = True,
    cold_outreach_limit: int = 10,
) -> dict:
    """
    Run the full pipeline and return structured JSON results.
    Called by api.py for each user-initiated run.
    The API key is used only for the duration of this call and never written to disk.
    """
    import anthropic
    client = anthropic.Anthropic(api_key=api_key)

    print(f"\n{'='*60}")
    print(f"  Pipeline run — {datetime.now().strftime('%Y-%m-%d %H:%M')} ({hours}h window)")
    print(f"{'='*60}\n")

    # --- Scraping ---
    all_jobs: list = []
    for label, fn, kw in [
        ("Greenhouse ATS",    scrape_greenhouse, {"max_h": hours}),
        ("Lever ATS",         scrape_lever,      {"max_h": hours}),
        ("Ashby/YC startups", scrape_ashby,      {"max_h": hours}),
        ("Wellfound",         scrape_wellfound,  {}),
        ("LinkedIn/Indeed",   scrape_jobspy,     {"max_h": hours}),
    ]:
        print(f"Scraping {label}...")
        found = fn(**kw)
        print(f"  {len(found)} internships found")
        all_jobs.extend(found)

    # --- Filtering ---
    PA_CITIES = ["philadelphia", "exton", "hatfield", "king of prussia", "malvern", "wayne", "conshohocken"]

    def _is_usa(loc):
        if not loc or not isinstance(loc, str):
            return False
        s = loc.lower()
        non_us = ["dublin", "luxembourg", "bengaluru", "toronto", "canada", "india", "brazil",
                  "ireland", "germany", "uk", "united kingdom", "france", "europe", "singapore",
                  "australia", "china", "japan", "hong kong", "mexico"]
        if any(x in s for x in non_us):
            return False
        if any(c in s for c in PA_CITIES) or re.search(r",?\s*pa(\s|$|,)", s):
            return True
        if re.fullmatch(r"\s*remote\s*", s) or re.fullmatch(r"\s*united states\s*", s):
            return True
        if "remote" in s and not any(x in s for x in non_us):
            return True
        if re.search(r"\busa\b|\bunited states\b|\bus\b", s) and not any(x in s for x in non_us):
            return True
        return False

    all_jobs = [j for j in all_jobs if is_intern(j["title"]) and _is_usa(j.get("location", ""))]

    seen: set = set()
    deduped: list = []
    for j in all_jobs:
        key = (j["company"].lower().strip(), j["title"].lower().strip())
        if key not in seen:
            seen.add(key)
            deduped.append(j)

    all_jobs = sorted(deduped, key=lambda x: x.get("hours_ago", 999))
    fresh = sum(1 for j in all_jobs if j.get("hours_ago", 999) <= 24)
    print(f"\nTotal: {len(all_jobs)} unique internships | {fresh} posted in last 24h\n")

    # --- Add score to each job ---
    for j in all_jobs:
        j["score"] = score(j["title"], j.get("description", ""))

    # --- Outreach (contact research + message drafting) ---
    outreach: list = []
    fresh_jobs = [j for j in all_jobs if j.get("hours_ago", 999) <= 96]
    print(f"Researching contacts for {min(len(fresh_jobs), 35)} jobs...\n")
    for j in fresh_jobs[:35]:
        co   = j.get("company", "")
        role = j.get("title", "")
        url  = j.get("job_url", "")
        if not co or not role:
            continue
        print(f"  {co} — {role}")
        contacts = find_contacts(co, role, client)
        for c in contacts[:2]:
            msgs = draft_messages(co, role, url, c, client)
            outreach.append({
                "company":     co,
                "role":        role,
                "name":        c.get("name", ""),
                "title":       c.get("title", ""),
                "reason":      c.get("reason", ""),
                "linkedin_url": c.get("linkedin_search_url", ""),
                "linkedin_msg": msgs.get("linkedin_message", ""),
                "email":       c.get("guessed_email", ""),
                "email_subj":  msgs.get("email_subject", ""),
                "email_body":  msgs.get("email_body", ""),
                "email_sent":  "",
            })

    # --- Cold Outreach ---
    cold_contacts: list = []
    if cold_outreach_enabled:
        try:
            from openpyxl import Workbook as _WB
            from cold_outreach import run_cold_outreach, build_config as _co_cfg
            co_config = _co_cfg()
            co_config["cold_outreach_max_per_day"] = cold_outreach_limit
            co_config["anthropic_api_key"] = api_key
            dummy_wb = _WB()
            cold_contacts = run_cold_outreach(dummy_wb, co_config, gmail_service=None)
        except Exception as e:
            print(f"WARNING: [Cold Outreach] {e}")

    timestamp = datetime.utcnow().isoformat() + "Z"
    return {
        "jobs":          all_jobs,
        "outreach":      outreach,
        "cold_outreach": cold_contacts,
        "summary": {
            "total_jobs":          len(all_jobs),
            "fresh_jobs":          fresh,
            "outreach_count":      len(outreach),
            "cold_outreach_count": len(cold_contacts),
        },
        "timestamp": timestamp,
    }


def main():
    p=argparse.ArgumentParser()
    p.add_argument("--no-email",action="store_true")
    p.add_argument("--hours",type=int,default=72)
    p.add_argument("--limit",type=int,default=None)
    p.add_argument("--scrape-only",action="store_true",
                   help="Run scrapers and write Excel only — skip all Claude API and Gmail calls")
    p.add_argument("--test-cold-outreach",action="store_true",
                   help="Dry-run the Cold Outreach module: scrapes Robin Hood only, "
                        "processes 2 companies, writes to 'Cold Outreach TEST' sheet, "
                        "does NOT update seen_companies.json")
    args=p.parse_args()

    # ----------------------------------------------------------------
    # --test-cold-outreach: standalone dry-run, exits after completion
    # ----------------------------------------------------------------
    if args.test_cold_outreach:
        from openpyxl import Workbook as _WB
        from cold_outreach import run_cold_outreach, build_config as _co_cfg
        print("\n" + "="*60)
        print("  Cold Outreach — TEST RUN")
        print("="*60)
        test_wb = _WB()
        co_config = _co_cfg()
        contacts = run_cold_outreach(
            test_wb, co_config,
            gmail_service=None,
            test_mode=True,
            test_source_limit=1,
            test_company_limit=2,
        )
        ts = datetime.now().strftime("%Y-%m-%d_%H-%M")
        test_fname = f"test_cold_outreach_{ts}.xlsx"
        test_wb.save(test_fname)
        print("\n" + "="*60)
        print(f"  TEST COMPLETE — {len(contacts)} contact(s) found")
        for c in contacts:
            print(f"  • {c.get('company_name','')} | {c.get('first_name','')} {c.get('last_name','')} "
                  f"| {c.get('email','(no email)')} | {c.get('email_status','')}")
        print(f"  Output: {test_fname}")
        print("  seen_companies.json was NOT modified (dry run)")
        print("="*60 + "\n")
        sys.exit(0)

    if not args.scrape_only:
        api_key=os.getenv("ANTHROPIC_API_KEY")
        if not api_key:
            print("Set ANTHROPIC_API_KEY in .env file or:\n  export ANTHROPIC_API_KEY='sk-...'")
            sys.exit(1)

        for var in ["YOUR_NAME", "YOUR_EMAIL", "YOUR_LINKEDIN", "YOUR_SCHOOL", "MY_BACKGROUND"]:
            if not os.getenv(var):
                print(f"WARNING: {var} is not set in .env — outreach messages may be incomplete")

        import anthropic
        client=anthropic.Anthropic(api_key=api_key)
    else:
        client=None
        api_key=None

    gmail=get_gmail() if (not args.no_email and not args.scrape_only) else None

    print(f"\n{'='*60}")
    print(f"  Job Pipeline — {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    print(f"  Finding jobs posted within {args.hours} hours")
    print(f"{'='*60}\n")

    all_jobs = []
    source_stats = {}  # Track raw vs filtered for each source
    
    for label, fn, kw in [
        ("Greenhouse ATS (startups/tech)", scrape_greenhouse, {"max_h": args.hours}),
        ("Lever ATS", scrape_lever, {"max_h": args.hours}),
        ("Ashby/YC startups", scrape_ashby, {"max_h": args.hours}),
        ("Wellfound/AngelList", scrape_wellfound, {}),
        ("LinkedIn/Indeed", scrape_jobspy, {"max_h": args.hours}),
    ]:
        print(f"Scraping {label}...")
        found = fn(**kw)
        raw_count = len(found)
        print(f"  {raw_count} internships found (raw)")
        all_jobs.extend(found)
        source_stats[label] = {"raw": raw_count, "filtered": 0}

    # Geographic filtering: Only keep jobs in PA, Remote, or US (no non-US cities/countries)
    PA_CITIES = [
        "philadelphia", "exton", "hatfield", "king of prussia", "malvern", "wayne", "conshohocken"
    ]
    def is_usa_location(loc):
        if not loc or not isinstance(loc, str):
            return False
        location_str = loc.lower()
        # Remove jobs with obvious non-US cities/countries
        non_us_keywords = [
            "dublin", "luxembourg", "bengaluru", "são paulo", "warsaw", "toronto", "canada", "india", "brazil", "ireland", "germany", "uk", "united kingdom", "france", "europe", "singapore", "australia", "china", "japan", "hong kong", "mexico", "spain", "netherlands", "sweden", "switzerland", "italy", "israel", "uae", "dubai", "new zealand", "south africa", "argentina", "chile", "colombia", "peru", "denmark", "norway", "finland", "belgium", "austria", "czech", "poland", "russia", "turkey", "egypt", "saudi", "korea", "taiwan", "malaysia", "thailand", "philippines", "indonesia", "vietnam", "romania", "portugal", "greece", "ukraine", "hungary", "slovakia", "slovenia", "croatia", "serbia", "bulgaria", "latvia", "lithuania", "estonia", "africa", "asia", "south america", "middle east"
        ]
        if any(x in location_str for x in non_us_keywords):
            return False
        # Keep if PA city or 'PA' in location
        if any(city in location_str for city in PA_CITIES) or re.search(r',?\s*pa(\s|$|,)', location_str):
            return True
        # Keep if Remote or United States (with no specific city)
        if re.fullmatch(r'\s*remote\s*', location_str) or re.fullmatch(r'\s*united states\s*', location_str):
            return True
        # If location is just 'Remote' or 'United States' (with no city), keep
        if location_str.strip() in ["remote", "united states"]:
            return True
        # If location contains 'remote' and no non-US city, keep
        if "remote" in location_str and not any(x in location_str for x in non_us_keywords):
            return True
        # If location contains 'usa' or 'us' and no non-US city, keep
        if re.search(r'\busa\b|\bunited states\b|\bus\b', location_str) and not any(x in location_str for x in non_us_keywords):
            return True
        return False

    all_jobs = [j for j in all_jobs if is_intern(j["title"]) and is_usa_location(j.get("location", ""))]
    
    # Deduplication: use compound key (company + title) to avoid duplicate roles at same company from different sources
    seen = set()
    deduped = []
    for j in all_jobs:
        # Create compound key: company name + role title (case-insensitive)
        company = j["company"].lower().strip()
        title = j["title"].lower().strip()
        dedup_key = (company, title)
        
        if dedup_key not in seen:
            seen.add(dedup_key)
            deduped.append(j)
    
    all_jobs = sorted(deduped, key=lambda x: x.get("hours_ago", 999))

    # Count filtered jobs by source
    for j in all_jobs:
        source = j.get("source", "Unknown")
        if source in source_stats:
            source_stats[source]["filtered"] += 1

    print("Per-Source Summary:")
    for source, stats in source_stats.items():
        raw = stats["raw"]
        filtered = stats["filtered"]
        pct = f"({filtered}/{raw})" if raw > 0 else "(0/0)"
        print(f"  {source}: {filtered} passing filters {pct}")
    print()

    if args.limit:
        all_jobs = all_jobs[:args.limit]

    fresh = sum(1 for j in all_jobs if j.get("hours_ago", 999) <= 24)
    print(f"\nTotal: {len(all_jobs)} unique internships | {fresh} posted in last 24h\n")

    # Filter for outreach: only jobs posted within 96 hours (4 days) — gives more time for contact research
    FRESHNESS_CUTOFF_H = 96
    fresh_jobs = [j for j in all_jobs if j.get("hours_ago", 999) <= FRESHNESS_CUTOFF_H]
    print(f"Jobs passing freshness check (≤{FRESHNESS_CUTOFF_H}h): {len(fresh_jobs)}/{len(all_jobs)}\n")

    outreach = []
    if args.scrape_only:
        print("--scrape-only: skipping Claude API calls and outreach drafting\n")
    else:
        print("Researching contacts and drafting outreach...\n")
        for j in fresh_jobs[:35]:
            co = j.get("company", "")
            role = j.get("title", "")
            url = j.get("job_url", "")
            if not co or not role:
                continue
            print(f"  {co} — {role} ({j.get('hours_ago', 999):.0f}h ago)")
            contacts = find_contacts(co, role, client)
            for c in contacts[:2]:
                msgs = draft_messages(co, role, url, c, client)
                email = c.get("guessed_email", "")
                sent = ""
                if email and not args.no_email:
                    delay = random.randint(EMAIL_DELAY_MIN, EMAIL_DELAY_MAX)
                    print(f"    Waiting {delay//60}m{delay%60}s before emailing {email}...")
                    time.sleep(delay)
                    ok = send_email(gmail, email, msgs.get("email_subject", ""), msgs.get("email_body", ""))
                    sent = "Sent" if ok else "Failed"
                elif email:
                    sent = "Drafted — send manually"
                else:
                    sent = "No email — LinkedIn only"
                outreach.append({
                    "company": co,
                    "role": role,
                    "name": c.get("name", ""),
                    "title": c.get("title", ""),
                    "reason": c.get("reason", ""),
                    "linkedin_url": c.get("linkedin_search_url", ""),
                    "linkedin_msg": msgs.get("linkedin_message", ""),
                    "email": email,
                    "email_subj": msgs.get("email_subject", ""),
                    "email_body": msgs.get("email_body", ""),
                    "email_sent": sent
                })

    fname, wb = write_excel(all_jobs, outreach)

    # ----------------------------------------------------------------
    # Cold Outreach module — runs after job scraping, adds its own tab
    # ----------------------------------------------------------------
    cold_contacts = []
    if not args.scrape_only:
        try:
            from cold_outreach import run_cold_outreach, build_config as _co_cfg
            co_config = _co_cfg()
            if co_config.get("cold_outreach_enabled", True):
                cold_contacts = run_cold_outreach(wb, co_config, gmail_service=gmail)
            else:
                print("[Cold Outreach] Disabled via COLD_OUTREACH_ENABLED=false")
        except Exception as e:
            print(f"WARNING: [Cold Outreach] Module error (main pipeline unaffected): {e}")

    # Save workbook (cold outreach sheet already appended above)
    wb.save(fname)

    print("\n" + "=" * 60)
    print("  COMPLETE")
    print(f"  {len(all_jobs)} jobs | {fresh} fresh | {len(outreach)} outreach drafted")
    if cold_contacts:
        print(f"  Cold Outreach: {len(cold_contacts)} new contacts added")
    print(f"  File: {fname}")
    print("\n  NEXT STEPS:")
    print("  1. Open file — start with GREEN rows (< 6h old)")
    print("  2. Go to Outreach tab — send LinkedIn messages (30 sec each)")
    print("  3. Run tomorrow morning: python job_pipeline_full.py --hours 24")
    print("=" * 60 + "\n")

if __name__=="__main__":
    main()
